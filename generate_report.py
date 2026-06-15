import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_report():
    doc = Document()
    
    # Title Page
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run("MERN STACK INTERNSHIP\nDAY 13 ASSIGNMENT REPORT\n(SHOPPING CART & CHECKOUT MODULE)")
    run.font.name = 'Arial'
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(107, 70, 193)  # Purple Accent
    
    # Subtitle / Info
    info_p = doc.add_paragraph()
    info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info_run = info_p.add_run(
        "\nAllocation Date – 15th June 2026\n"
        "Submission Date – 15th June 2026\n\n"
        "Name: PARTH SHARMA\n"
        "Role: MERN Stack Developer Intern\n"
        "Position: MERN Stack Intern\n"
        "Internship Organization: Insta Dot Analytics\n"
    )
    info_run.font.name = 'Arial'
    info_run.font.size = Pt(12)
    
    doc.add_page_break()
    
    # Section 1: Objective & Description
    h1 = doc.add_paragraph()
    r1 = h1.add_run("TASK 1 - SHOPPING CART SCHEMA AND MONGODB STORAGE")
    r1.font.name = 'Arial'
    r1.font.size = Pt(16)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(107, 70, 193)
    
    p = doc.add_paragraph()
    p.add_run("Objective:\n").bold = True
    p.add_run("To design and implement a persistent Shopping Cart data model in MongoDB linked to authenticated users, supporting product aggregation, quantity increments, and item removals.\n\n")
    p.add_run("Description:\n").bold = True
    p.add_run("A MongoDB schema for the Shopping Cart was created using Mongoose. The model maintains a reference to the User schema and holds an array of cart items, each storing a reference to the Product model and the selected quantity.\n\n")
    p.add_run("Implementation:\n").bold = True
    p.add_run("- Built the Cart Mongoose Schema with validation (minimum quantity of 1).\n- Configured relational populates to resolve full product details dynamically.\n- Associated cart items uniquely per authenticated user.\n")
    
    # Code snippet
    p_code_header = doc.add_paragraph()
    p_code_header.add_run("Mongoose Cart Schema Model (backend/models/Cart.js):").italic = True
    p_code = doc.add_paragraph()
    p_code_run = p_code.add_run(
        "import mongoose from 'mongoose';\n\n"
        "const CartItemSchema = new mongoose.Schema({\n"
        "  product: {\n"
        "    type: mongoose.Schema.Types.ObjectId,\n"
        "    ref: 'Product',\n"
        "    required: true,\n"
        "  },\n"
        "  quantity: {\n"
        "    type: Number,\n"
        "    required: true,\n"
        "    min: [1, 'Quantity must be at least 1'],\n"
        "    default: 1,\n"
        "  },\n"
        "});\n\n"
        "const CartSchema = new mongoose.Schema({\n"
        "  user: { type: mongoose.Schema.Types.ObjectId, ref: 'User', required: true, unique: true },\n"
        "  items: [CartItemSchema],\n"
        "}, { timestamps: true });\n\n"
        "export default mongoose.model('Cart', CartSchema);"
    )
    p_code_run.font.name = 'Courier New'
    p_code_run.font.size = Pt(9.5)
    
    doc.add_page_break()
    
    # Section 2: Cart Controller and REST APIs
    h2 = doc.add_paragraph()
    r2 = h2.add_run("TASK 2 - CART CRUD REST APIs (EXPRESS.JS)")
    r2.font.name = 'Arial'
    r2.font.size = Pt(16)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(107, 70, 193)
    
    p = doc.add_paragraph()
    p.add_run("Objective:\n").bold = True
    p.add_run("To construct REST endpoints allowing client-side actions to fetch, add, update, and remove items from the database-backed cart.\n\n")
    p.add_run("Description:\n").bold = True
    p.add_run("Using Express routers, endpoints were registered and protected by JWT auth middleware. The controller handles add operations by incrementing existing item counts or pushing new array elements, updating quantities, and clearing carts.\n\n")
    p.add_run("Endpoints Implemented:\n").bold = True
    p.add_run("- GET /api/cart - Fetches authenticated user's populated cart.\n")
    p.add_run("- POST /api/cart/add - Adds item or increases quantity.\n")
    p.add_run("- PUT /api/cart/update - Modifies quantity of a target product.\n")
    p.add_run("- DELETE /api/cart/remove/:productId - Deletes product from the cart.\n")
    p.add_run("- POST /api/cart/clear - Truncates the items array.\n")
    
    # Section 3: React Hooks for State Management
    h3 = doc.add_paragraph()
    r3 = h3.add_run("TASK 3 - REACT HOOKS AND CART CONTEXT STATE MANAGEMENT")
    r3.font.name = 'Arial'
    r3.font.size = Pt(16)
    r3.font.bold = True
    r3.font.color.rgb = RGBColor(107, 70, 193)
    
    p = doc.add_paragraph()
    p.add_run("Objective:\n").bold = True
    p.add_run("To establish a centralized React Context with custom hooks (`useCart`) to consume cart operations reactively across catalog, cart details, and checkout views.\n\n")
    p.add_run("Description:\n").bold = True
    p.add_run("The `CartProvider` tracks local cart state using React hooks and updates immediately upon user interactions. Requests are dispatched asynchronously to the MERN backend using authenticated fetch calls.\n")
    
    doc.add_page_break()
    
    # Section 4: Responsive UI
    h4 = doc.add_paragraph()
    r4 = h4.add_run("TASK 4 - RESPONSIVE SHOPPING CART & CHECKOUT PAGE UI")
    r4.font.name = 'Arial'
    r4.font.size = Pt(16)
    r4.font.bold = True
    r4.font.color.rgb = RGBColor(107, 70, 193)
    
    p = doc.add_paragraph()
    p.add_run("Objective:\n").bold = True
    p.add_run("To develop state-of-the-art, responsive web views for the Cart and Checkout operations matching custom glassmorphism styling parameters.\n\n")
    p.add_run("Features:\n").bold = True
    p.add_run("- Cart Page: Dynamic unit price and subtotal calculation, quantity selectors, checkout redirection, and free shipping calculations.\n- Checkout Page: Complete shipping address form, mock Credit/Debit Card payment forms, responsive order summaries, and success completion cards.\n")
    
    # Insert screenshot
    img_path = "/Users/macbook/.gemini/antigravity-ide/brain/42d8c101-f008-42a7-ab63-e9b9808302d2/day13_shopping_cart_checkout_1781525841608.png"
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_img = p_img.add_run()
        r_img.add_picture(img_path, width=Inches(5.5))
        p_caption = doc.add_paragraph()
        p_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rc = p_caption.add_run("Figure 13.1: Premium Responsive Cart and Checkout Interface Mockup")
        rc.font.italic = True
        rc.font.size = Pt(10)
        
    doc.add_page_break()
    
    # Final Conclusion
    h5 = doc.add_paragraph()
    r5 = h5.add_run("FINAL CONCLUSION")
    r5.font.name = 'Arial'
    r5.font.size = Pt(16)
    r5.font.bold = True
    r5.font.color.rgb = RGBColor(107, 70, 193)
    
    p = doc.add_paragraph()
    p.add_run("The addition of the Shopping Cart and Checkout pages concludes a complete transaction flow for the MERN Stack application. Persistent database storage in MongoDB guarantees cart stability between user sessions, and custom React Hooks provide seamless client state updates. The responsive layout provides a stellar mobile-first browsing experience.\n\n")
    p.add_run("GitHub Code Repository: ").bold = True
    p.add_run("https://github.com/parthsharma17prs/day13-shopping-cart\n\n")
    p.add_run("Thank you!\n")
    
    doc.save("Parth_Sharma_MERN_Day13_Report.docx")
    print("Report generated successfully as Parth_Sharma_MERN_Day13_Report.docx")

if __name__ == '__main__':
    create_report()
